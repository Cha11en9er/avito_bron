from __future__ import annotations

import base64
import io
import json
import os
import random
import re
import sys
import traceback
from datetime import datetime
from pathlib import Path

import ddddocr
from PIL import Image, ImageOps
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from playwright_stealth.stealth import Stealth

NAV_TIMEOUT_MS = 90000
ON_PAGE_BASE_WAIT_S = 8
BETWEEN_SITES_PAUSE_S = 30
URLS_FILE_NAME = "urls.txt"
OUTPUT_DIR_NAME = "avito_houses_dump"
RESULTS_FILE_NAME = "avito_houses_results.json"
RESULTS_XLSX_FILE_NAME = "дома авито.xlsx"
NOT_FOUND_VALUE = "нету на сайте"
PHONE_NOT_RECOGNIZED = "не распознан"

_OCR = ddddocr.DdddOcr(show_ad=False)

EXPORT_COLUMNS = [
    "номер",
    "название",
    "цена",
    "адрес",
    "телефон",
    "площадь дома",
    "площадь участка",
    "этажей",
    "материал стен",
    "год постройки",
    "описание",
    "характеристики",
    "ссылка",
]

HOUSE_FIELD_KEYWORDS: dict[str, list[str]] = {
    "площадь дома":    ["площадь дома", "общая площадь"],
    "площадь участка": ["площадь участка"],
    "этажей":          ["этажей в доме", "этажност", "этажей"],
    "материал стен":   ["материал стен", "тип дома"],
    "год постройки":   ["год постройки"],
}

CHROME_ARGS = [
    "--start-maximized",
    "--disable-blink-features=AutomationControlled",
    "--disable-dev-shm-usage",
    "--no-first-run",
    "--disable-features=IsolateOrigins,site-per-process",
]

HUMAN_USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/133.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/132.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
]


def _setup_playwright_env() -> None:
    if getattr(sys, "frozen", False):
        base = Path(sys.executable).resolve().parent
    else:
        base = Path(__file__).resolve().parent
    bundled = base / "ms-playwright"
    if bundled.is_dir():
        os.environ["PLAYWRIGHT_BROWSERS_PATH"] = str(bundled)


def _console_utf8() -> None:
    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
            sys.stderr.reconfigure(encoding="utf-8")
        except Exception:
            pass


def _launch_browser(playwright):
    try:
        return playwright.chromium.launch(channel="chrome", headless=False, args=CHROME_ARGS)
    except Exception:
        return playwright.chromium.launch(headless=False, args=CHROME_ARGS)


def _new_page(browser):
    user_agent = random.choice(HUMAN_USER_AGENTS)
    viewport_w = random.choice([1920, 1912, 1904])
    viewport_h = random.choice([1080, 1032, 1040, 1008])
    context = browser.new_context(
        viewport={"width": viewport_w, "height": viewport_h},
        user_agent=user_agent,
        locale="ru-RU",
        timezone_id="Europe/Moscow",
        color_scheme="light",
        device_scale_factor=1,
    )
    try:
        stealth = Stealth(
            navigator_languages_override=("ru-RU", "ru", "en-US", "en"),
            navigator_platform_override="Win32",
            webgl_vendor_override="Intel Inc.",
            webgl_renderer_override="Intel Iris OpenGL Engine",
            navigator_user_agent_override=user_agent,
        )
    except TypeError:
        # Совместимость со старыми версиями playwright-stealth.
        stealth = Stealth(
            navigator_languages_override=("ru-RU", "ru", "en-US", "en"),
            navigator_platform_override="Win32",
            webgl_vendor_override="Intel Inc.",
            webgl_renderer_override="Intel Iris OpenGL Engine",
        )
    stealth.apply_stealth_sync(context)
    context.set_extra_http_headers(
        {
            "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
            "Upgrade-Insecure-Requests": "1",
            "Sec-CH-UA-Platform": '"Windows"',
            "Sec-CH-UA-Mobile": "?0",
        }
    )
    page = context.new_page()
    page.set_default_navigation_timeout(NAV_TIMEOUT_MS)
    page.set_default_timeout(25000)
    return context, page


def _human_pause(page, min_seconds: float = 0.4, max_seconds: float = 1.6) -> None:
    wait_ms = int(random.uniform(min_seconds, max_seconds) * 1000)
    page.wait_for_timeout(wait_ms)


def _simulate_human_activity(page) -> None:
    # Лёгкие случайные движения мышью — убираем "идеально ровный" машинный паттерн.
    try:
        for _ in range(random.randint(2, 4)):
            x = random.randint(100, 1200)
            y = random.randint(120, 760)
            page.mouse.move(x, y, steps=random.randint(8, 20))
            _human_pause(page, 0.1, 0.35)
    except Exception:
        pass


def _scroll_through_page(page, total_scrolls: int = 6) -> None:
    """Прокручиваем страницу постепенно: вниз пачкой, иногда чуть обратно, в конце — наверх."""
    try:
        size = page.viewport_size or {"height": 900}
        viewport_h = size.get("height") or 900
        for _ in range(total_scrolls):
            step = int(viewport_h * random.uniform(0.45, 0.9))
            page.mouse.wheel(0, step)
            _human_pause(page, 0.5, 1.3)
            if random.random() < 0.3:
                page.mouse.wheel(0, -int(step * random.uniform(0.2, 0.45)))
                _human_pause(page, 0.3, 0.8)
        page.evaluate("window.scrollTo({ top: 0, behavior: 'smooth' })")
        _human_pause(page, 0.8, 1.6)
    except Exception:
        pass


def _safe_slug(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() else "_" for ch in value.strip().lower())
    cleaned = "_".join(part for part in cleaned.split("_") if part)
    return cleaned[:80] or "house"


def _normalize_space(value: str) -> str:
    return " ".join((value or "").split()).strip()


def _load_urls(base_dir: Path) -> list[str]:
    urls_path = base_dir / URLS_FILE_NAME
    if not urls_path.exists():
        raise FileNotFoundError(
            f"Не найден файл со ссылками: {urls_path}. Создайте {URLS_FILE_NAME} рядом со скриптом."
        )
    urls: list[str] = []
    seen: set[str] = set()
    with open(urls_path, "r", encoding="utf-8") as f:
        for raw in f:
            line = raw.strip()
            if not line or line.startswith("#"):
                continue
            if line in seen:
                continue
            seen.add(line)
            urls.append(line)
    if not urls:
        raise ValueError(f"{URLS_FILE_NAME} не содержит ни одной валидной ссылки")
    return urls


def _extract_id_from_url(url: str) -> int | None:
    cleaned = url.split("?", 1)[0].rstrip("/")
    match = re.search(r"_(\d{6,})$", cleaned)
    if match:
        return int(match.group(1))
    return None


def _load_existing_results(base_dir: Path) -> list[dict]:
    results_path = base_dir / RESULTS_FILE_NAME
    if not results_path.exists():
        return []
    try:
        with open(results_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except Exception:
        return []


def _write_results(base_dir: Path, payload: list[dict]) -> None:
    results_path = base_dir / RESULTS_FILE_NAME
    with open(results_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _write_excel(base_dir: Path, payload: list[dict]) -> None:
    out_path = base_dir / RESULTS_XLSX_FILE_NAME
    wb = Workbook()
    ws = wb.active
    ws.title = "Дома"
    ws.append(EXPORT_COLUMNS)
    for row in payload:
        ws.append([row.get(col, "") for col in EXPORT_COLUMNS])
    wb.save(out_path)


def _is_record_quality_good(record: dict) -> bool:
    """
    "Хорошая" запись = хотя бы одно из ключевых полей реально извлечено.
    Заглушки `нету на сайте` / `не распознан` считаем признаком капчи/ошибки.
    """
    link = _normalize_space(str(record.get("ссылка", "")))
    if not link:
        return False
    placeholders = {NOT_FOUND_VALUE, PHONE_NOT_RECOGNIZED, ""}
    key_fields = [
        "название",
        "цена",
        "адрес",
        "телефон",
        "площадь дома",
        "площадь участка",
        "описание",
    ]
    for field in key_fields:
        val = _normalize_space(str(record.get(field, "")))
        if val and val not in placeholders:
            return True
    return False


def _extract_title(page) -> str:
    for selector in (
        'h1[itemprop="name"]',
        '[data-marker="item-view/title-info"] h1',
        "h1",
    ):
        loc = page.locator(selector).first
        if loc.count() > 0:
            text = _normalize_space(loc.inner_text())
            if text:
                return text
    return ""


def _extract_price(page) -> str:
    candidates = (
        '[data-marker="item-view/item-price"]',
        'span[data-marker="item-view/item-price-content"]',
        '[itemprop="price"]',
    )
    for selector in candidates:
        loc = page.locator(selector).first
        if loc.count() == 0:
            continue
        try:
            content_attr = loc.get_attribute("content")
            if content_attr and re.search(r"\d", content_attr):
                return _normalize_space(content_attr)
        except Exception:
            pass
        text = _normalize_space(loc.inner_text())
        if text:
            return text
    return ""


def _extract_address(page) -> str:
    for selector in (
        '[itemprop="address"] span._8360df6eedcf8d52',
        '[itemprop="address"]',
        'div[class*="style-item-address"]',
    ):
        loc = page.locator(selector).first
        if loc.count() > 0:
            text = _normalize_space(loc.inner_text())
            if text:
                return text
    return ""


def _extract_description(page) -> str:
    loc = page.locator('div[data-marker="item-view/item-description"]').first
    if loc.count() == 0:
        return ""
    return _normalize_space(loc.inner_text())


def _extract_details_map(page) -> dict[str, str]:
    """Достаём пары `параметр: значение` из блока характеристик."""
    details: dict[str, str] = {}

    rows = page.locator("li.d2936d013c910379")
    if rows.count() == 0:
        # Резервный селектор — на случай, если CSS-хеши Авито поменялись.
        rows = page.locator(
            '[data-marker="item-view/params"] li, '
            'ul[class*="params"] li, '
            'div[class*="params"] li'
        )

    count = rows.count()
    for i in range(count):
        row = rows.nth(i)
        label_locator = row.locator(
            'span.d6e8fd2e3d52b32a, span[class*="styles-module-label"], span[class*="param-label"]'
        ).first
        full_text = _normalize_space(row.inner_text())
        label = ""
        value = ""
        if label_locator.count() > 0:
            label_text = _normalize_space(label_locator.inner_text())
            label = label_text.rstrip(":").lower()
            value = full_text.replace(label_text, "", 1).strip(" :")
        elif ":" in full_text:
            head, tail = full_text.split(":", 1)
            label = _normalize_space(head).lower()
            value = _normalize_space(tail)
        if label:
            details[label] = _normalize_space(value)
    return details


def _pick_house_field(details: dict[str, str], target_key: str) -> str:
    for label, value in details.items():
        for kw in HOUSE_FIELD_KEYWORDS.get(target_key, [target_key]):
            if kw in label:
                return value
    return ""


def _save_phone_image_from_popup(
    page,
    save_dir: Path,
    item_id: int,
    title: str,
    timestamp: str,
) -> str | None:
    button = page.locator('button[data-marker="item-phone-button/card"]').first
    if button.count() == 0:
        return None
    button.click(timeout=10000)

    popup = page.locator("div.ea88242d5926e753.b0253be71a159c76").first
    popup.wait_for(state="visible", timeout=10000)

    image = popup.locator("img._71e723cdca9c6624").first
    image.wait_for(state="visible", timeout=10000)
    src = image.get_attribute("src")
    if not src:
        return None

    content: bytes
    extension = "png"

    if src.startswith("data:image/"):
        header, _, encoded = src.partition(",")
        match = re.match(r"^data:image/([a-zA-Z0-9.+-]+);base64$", header)
        if not match:
            return None
        extension = match.group(1).replace("+xml", "")
        content = base64.b64decode(encoded)
    else:
        if src.startswith("//"):
            src = f"https:{src}"
        response = page.request.get(src, timeout=15000)
        if not response.ok:
            return None
        content_type = (response.header_value("content-type") or "").lower()
        match = re.search(r"image/([a-z0-9.+-]+)", content_type)
        if match:
            extension = match.group(1).replace("+xml", "")
        content = response.body()

    file_name = f"{item_id}_{_safe_slug(title)}_{timestamp}_phone.{extension}"
    out_path = save_dir / file_name
    with open(out_path, "wb") as f:
        f.write(content)
    return str(out_path)


def _extract_russian_phone(text: str) -> str | None:
    prepared = (
        (text or "")
        .replace("O", "0")
        .replace("o", "0")
        .replace("I", "1")
        .replace("l", "1")
        .replace("S", "5")
        .replace("B", "8")
    )
    cleaned = re.sub(r"[^\d+]", "", prepared)
    candidates: list[str] = []
    if cleaned:
        candidates.append(cleaned)
    candidates.extend(re.findall(r"(?:\+?7|8)\D*\d\D*\d\D*\d\D*\d\D*\d\D*\d\D*\d\D*\d\D*\d\D*\d", text))

    for candidate in candidates:
        digits = re.sub(r"\D", "", candidate)
        if len(digits) == 11 and (digits.startswith("8") or digits.startswith("7")):
            if digits.startswith("7"):
                digits = "8" + digits[1:]
            return digits
    return None


def _ocr_phone_number(image_path: str) -> str | None:
    with Image.open(image_path) as source_image:
        rgba = source_image.convert("RGBA")
        white_bg = Image.new("RGBA", rgba.size, "WHITE")
        white_bg.alpha_composite(rgba)
        image = ImageOps.grayscale(white_bg)
        # Увеличиваем и бинаризуем: для крупных чёрных цифр на белом фоне сильно повышает OCR-точность.
        image = image.resize((image.width * 3, image.height * 3))
        image = image.point(lambda p: 255 if p > 180 else 0)

    buf = io.BytesIO()
    image.save(buf, format="PNG")
    ocr_text = _OCR.classification(buf.getvalue())
    return _extract_russian_phone(ocr_text)


def _or_not_found(value: str | None) -> str:
    normalized = _normalize_space(value or "")
    return normalized if normalized else NOT_FOUND_VALUE


def _process_one_url(
    page,
    output_dir: Path,
    base_dir: Path,
    url: str,
    idx: int,
    total: int,
) -> dict | None:
    item_id = _extract_id_from_url(url) or idx
    print(f"[{idx}/{total}] Открываю: {url}")
    page.goto(url, wait_until="domcontentloaded", timeout=NAV_TIMEOUT_MS)
    _human_pause(page, 1.5, 4.0)
    _simulate_human_activity(page)
    dynamic_wait = ON_PAGE_BASE_WAIT_S + random.randint(2, 8)
    print(f"[{idx}/{total}] Жду {dynamic_wait} сек., затем прокручиваю страницу...")
    page.wait_for_timeout(dynamic_wait * 1000)
    _scroll_through_page(page, total_scrolls=random.randint(5, 8))
    _simulate_human_activity(page)
    _human_pause(page, 1.0, 2.5)

    title = _extract_title(page) or f"house_{item_id}"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = f"{item_id}_{_safe_slug(title)}_{timestamp}.html"
    html_path = output_dir / file_name
    try:
        html_content = page.content()
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html_content)
    except Exception as exc:
        print(f"[{idx}/{total}] id={item_id} Не удалось сохранить HTML: {exc}")

    try:
        phone_image_path = _save_phone_image_from_popup(
            page=page,
            save_dir=base_dir,
            item_id=item_id,
            title=title,
            timestamp=timestamp,
        )
    except Exception as exc:
        print(f"[{idx}/{total}] id={item_id} Не удалось сохранить картинку телефона: {exc}")
        phone_image_path = None

    phone_number: str | None = None
    if phone_image_path:
        try:
            phone_number = _ocr_phone_number(phone_image_path)
        except Exception as exc:
            print(f"[{idx}/{total}] id={item_id} Ошибка OCR телефона: {exc}")
            phone_number = None
    if phone_image_path and phone_number:
        try:
            Path(phone_image_path).unlink(missing_ok=True)
            phone_image_path = None
        except Exception as exc:
            print(f"[{idx}/{total}] id={item_id} Не удалось удалить картинку телефона: {exc}")

    price = _extract_price(page)
    address = _extract_address(page)
    description = _extract_description(page)
    details = _extract_details_map(page)

    record = {
        "номер": item_id,
        "название": _or_not_found(title),
        "цена": _or_not_found(price),
        "адрес": _or_not_found(address),
        "телефон": phone_number if phone_number else (PHONE_NOT_RECOGNIZED if phone_image_path else NOT_FOUND_VALUE),
        "площадь дома": _or_not_found(_pick_house_field(details, "площадь дома")),
        "площадь участка": _or_not_found(_pick_house_field(details, "площадь участка")),
        "этажей": _or_not_found(_pick_house_field(details, "этажей")),
        "материал стен": _or_not_found(_pick_house_field(details, "материал стен")),
        "год постройки": _or_not_found(_pick_house_field(details, "год постройки")),
        "описание": _or_not_found(description),
        "характеристики": _or_not_found(json.dumps(details, ensure_ascii=False)) if details else NOT_FOUND_VALUE,
        "ссылка": url,
    }
    print(f"[{idx}/{total}] Сохранено: {html_path.name}")
    return record


def main() -> None:
    _setup_playwright_env()

    base_dir = Path(__file__).resolve().parent
    output_dir = base_dir / OUTPUT_DIR_NAME
    output_dir.mkdir(parents=True, exist_ok=True)

    urls = _load_urls(base_dir)
    print(f"К обработке {len(urls)} ссылок.")

    results: list[dict] = _load_existing_results(base_dir)
    processed_links_ok = {
        _normalize_space(str(item.get("ссылка", "")))
        for item in results
        if isinstance(item, dict) and _is_record_quality_good(item)
    }

    with sync_playwright() as p:
        browser = _launch_browser(p)
        context, page = _new_page(browser)
        try:
            total = len(urls)
            first_visit = True
            for idx, url in enumerate(urls, start=1):
                if url in processed_links_ok:
                    print(f"[{idx}/{total}] Уже обработано хорошо, пропускаю: {url}")
                    continue

                if not first_visit:
                    pause_s = BETWEEN_SITES_PAUSE_S + random.randint(0, 10)
                    print(f"[{idx}/{total}] Пауза между сайтами: {pause_s} сек.")
                    page.wait_for_timeout(pause_s * 1000)
                first_visit = False

                try:
                    record = _process_one_url(page, output_dir, base_dir, url, idx, total)
                except Exception as exc:
                    print(f"[{idx}/{total}] Ошибка обработки {url}: {exc}")
                    traceback.print_exc()
                    record = None

                if record is None:
                    continue

                replaced = False
                for i, existing in enumerate(results):
                    if (
                        isinstance(existing, dict)
                        and _normalize_space(str(existing.get("ссылка", ""))) == url
                    ):
                        results[i] = record
                        replaced = True
                        break
                if not replaced:
                    results.append(record)

                if _is_record_quality_good(record):
                    processed_links_ok.add(url)

                _write_results(base_dir, results)
                _write_excel(base_dir, results)

            print(f"JSON сохранён: {base_dir / RESULTS_FILE_NAME}")
            print(f"Excel сохранён: {base_dir / RESULTS_XLSX_FILE_NAME}")
        finally:
            try:
                context.close()
            except Exception:
                pass
            browser.close()


if __name__ == "__main__":
    _console_utf8()
    try:
        main()
    except Exception:
        print("Произошла ошибка:")
        traceback.print_exc()
        sys.exit(1)
